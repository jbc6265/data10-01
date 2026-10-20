"""기획서 Markdown → Word(docx): python3 scripts/make-docx.py
docs/01_프로젝트_기획서.md 를 읽어 같은 이름의 .docx 를 만듭니다(python-docx 필요).
제목(#)·표(|)·글머리(-)·번호(1.)·코드 블록(```)·**굵게** 만 다룹니다. 코드 블록(mermaid 도해)은 고정폭 글자로 그대로 넣습니다."""
import re
import sys
from pathlib import Path
from docx import Document
from docx.enum.table import WD_TABLE_ALIGNMENT
from docx.oxml.ns import qn
from docx.shared import Pt, RGBColor

ROOT = Path(__file__).resolve().parent.parent
SRC = ROOT / 'docs' / '01_프로젝트_기획서.md'
OUT = SRC.with_suffix('.docx')
FONT = '맑은 고딕'


def set_font(run, size=None, mono=False):
    name = 'Consolas' if mono else FONT
    run.font.name = name
    run._element.rPr.rFonts.set(qn('w:eastAsia'), '맑은 고딕' if not mono else 'D2Coding')
    if size:
        run.font.size = Pt(size)


def add_inline(par, text, size=None):
    # **굵게** 와 `코드` 만 나눠 넣습니다
    for part in re.split(r'(\*\*[^*]+\*\*|`[^`]+`)', text):
        if not part:
            continue
        if part.startswith('**') and part.endswith('**'):
            r = par.add_run(part[2:-2]); r.bold = True; set_font(r, size)
        elif part.startswith('`') and part.endswith('`'):
            r = par.add_run(part[1:-1]); set_font(r, size, mono=True)
        else:
            r = par.add_run(part); set_font(r, size)


def table(doc, rows):
    cells = [[c.strip() for c in r.strip().strip('|').split('|')] for r in rows]
    cells = [r for r in cells if not all(re.fullmatch(r':?-{3,}:?', c) for c in r)]
    t = doc.add_table(rows=len(cells), cols=max(len(r) for r in cells))
    t.style = 'Table Grid'
    t.alignment = WD_TABLE_ALIGNMENT.CENTER
    for i, r in enumerate(cells):
        for j, c in enumerate(r):
            p = t.cell(i, j).paragraphs[0]
            add_inline(p, c.replace('<br/>', ' '), 9)
            if i == 0:
                for run in p.runs:
                    run.bold = True
    doc.add_paragraph()


def main():
    doc = Document()
    st = doc.styles['Normal']; st.font.name = FONT; st.font.size = Pt(10.5)
    st.element.rPr.rFonts.set(qn('w:eastAsia'), FONT)
    lines = SRC.read_text(encoding='utf-8').splitlines()
    i = 0
    while i < len(lines):
        ln = lines[i]
        if ln.startswith('```'):
            j = i + 1
            while j < len(lines) and not lines[j].startswith('```'):
                j += 1
            p = doc.add_paragraph()
            r = p.add_run('\n'.join(lines[i + 1:j])); set_font(r, 8.5, mono=True)
            r.font.color.rgb = RGBColor(0x33, 0x3d, 0x4a)
            i = j + 1; continue
        if ln.startswith('|'):
            j = i
            while j < len(lines) and lines[j].startswith('|'):
                j += 1
            table(doc, lines[i:j]); i = j; continue
        m = re.match(r'^(#{1,4})\s+(.*)', ln)
        if m:
            hd = doc.add_heading(level=min(len(m.group(1)) - 1, 3) if len(m.group(1)) > 1 else 0)
            add_inline(hd, m.group(2))
        elif re.match(r'^\s*-\s+', ln):
            p = doc.add_paragraph(style='List Bullet'); add_inline(p, re.sub(r'^\s*-\s+', '', ln))
        elif re.match(r'^\d+\.\s+', ln):
            p = doc.add_paragraph(style='List Number'); add_inline(p, re.sub(r'^\d+\.\s+', '', ln))
        elif ln.strip() == '---' or not ln.strip():
            pass
        else:
            add_inline(doc.add_paragraph(), ln)
        i += 1
    doc.save(OUT)
    print('만듦:', OUT.relative_to(ROOT))


if __name__ == '__main__':
    sys.exit(main())
